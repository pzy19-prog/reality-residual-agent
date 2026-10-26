#!/usr/bin/env python3
"""Build the complete, portable RRA submission bundle from adjacent data/ files.

Usage: python build_submission.py --output-dir /path/to/SUBMISSION_FINAL
No repository-root, control-directory, network, model, or evaluation access is used.
"""
from __future__ import annotations
import argparse, hashlib, html, json, math, shutil, subprocess, textwrap, wave, zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import fitz
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

HERE=Path(__file__).resolve().parent
DATA=HERE/"data"
FACTS=json.loads((DATA/"formal_facts.json").read_text())
TRACE=json.loads((DATA/"dev_trace.json").read_text())
PROVENANCE=json.loads((DATA/"source_provenance.json").read_text())
W,H=13.333,7.5
NAVY="101B2B"; PANEL="17283B"; PANEL2="243349"; TEAL="35D0C3"; AMBER="FFB454"; WHITE="F4F7FB"; MUTED="A9B9C8"; RED="FF7E76"
FONT="/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B="/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
pdfmetrics.registerFont(TTFont('RRA-DVSans',FONT))
pdfmetrics.registerFont(TTFont('RRA-DVSans-Bold',FONT_B))
SLIDES=[]; VIDEO_SCENES=[]; OUT:Path

def run(args): subprocess.run(args,check=True)
def rgb(s): return tuple(int(s[i:i+2],16) for i in (0,2,4))
def fmt(x,n=3): return f"{x:.{n}f}"
def shortsha(s): return s[:12]

def source_fields():
    r=FACTS["formal"]
    return r["b2"],r["b3"],r["b4"]

def make_slide_specs():
    b2,b3,b4=source_fields(); trace=TRACE
    return [
      dict(k="01 / REALITY RESIDUAL AGENT",title="Recovery decisions under uncertainty",sub="A 2D conveyor simulation and evidence-constrained recovery agent",kind="cover",body=["For automation and robotics experimentation", "Track 03 · Reasoning Architecture", "A bounded research prototype — not a deployed robot"]),
      dict(k="01 / THE PROBLEM",title="A mismatch is not a diagnosis",kind="problem",body=["Sensor bias", "Velocity offset", "Model error", "One symptom can have multiple causes. Recovery begins with evidence."]),
      dict(k="02 / THE APPROACH",title="Propose, verify, then authorize",kind="architecture",body=["FAST RESIDUAL", "B3 EXHAUSTIVE DSL", "B4 NEMOTRON PROPOSES / RANKS", "DETERMINISTIC FIT + COUNTERFACTUAL VERIFY", "COMMON-ACTION GATE → BOUNDED AUTHORIZATION → RECEIPT", "A language model does not control actuators or replace verification."]),
      dict(k="03 / PRODUCT VIEW",title="Real decisions from a development trace",kind="product",body=[f"{trace['episode_id']} · {len(trace['decisions'])} persisted decisions", "Development scenario: SENSOR_STUCK", "Four shown checks pass; common safe action updates", "AMBIGUOUS: no common safe action across shown passers", "Policy fault label: BELT_STOP_OR_MOTION_LOSS", f"{trace['terminal_outcomes']['picked']} picked · {trace['terminal_outcomes']['escalated']} escalated · replay match", "Submission-rendered data; not live or formal."]),
      dict(k="04 / ZETARIS",title="Federate governed evidence sources",kind="zetaris",body=["PostgreSQL telemetry", "Read-only REST sources", "Zetaris federation", "Whitelisted C1 evidence packet", "Recorded local-reference equivalence: T3/T4 mismatches 0 · T8 bytes identical"]),
      dict(k="05 / METERLESS",title="Persist and replay the reasoning trail",kind="meterless",body=["Episode run + manifest", "Pinned local H-MEM reference", "Episode-scoped read / cross-episode refusal", "Manifest replay → decision and receipt comparison", f"Fresh development checks: {FACTS['meterless_dev_fresh']['episodes']} episodes · replay {FACTS['meterless_dev_fresh']['M4_replay_match']}/{FACTS['meterless_dev_fresh']['episodes']} · cross-episode reads {FACTS['meterless_dev_fresh']['M3_cross_episode_reads']}"]),
      dict(k="06 / FORMAL COMPARISON",title="TIEBREAK showed no recovery gain",kind="chart",body=[]),
      dict(k="07 / INTERPRETATION",title="Keep counts and accounting units distinct",kind="interpret",body=[]),
      dict(k="08 / NVIDIA + STACK",title="Nemotron proposes; deterministic gates decide",kind="sponsors",body=["NVIDIA Nemotron: B4 candidate structure proposal and ranking.", "Fit, counterfactual verification, common-action gate and bounded authorization remain deterministic.", "Python 3.12 · NumPy 2.5.4 · local Python CLI · Meterless TypeScript / H-MEM", "No measured sponsor-caused recovery gain is claimed."]),
      dict(k="09 / LIMITS + DISCLOSURE",title="What the evidence does not establish",kind="limits",body=["B3 / STRICT: 15 unsafe fast-layer events; S2-authorized unsafe events: 0. B3 had 159 authorizations.", "Case 9 was seen in development; no unseen-composition generalization claim.", "2D simulation only: no physical deployment, production customers, measured ROI or broad safe-agent guarantee.", "Pre-window boundary includes V0 plus H1 runtime, eval/evidence foundation and preregistration."]),
      dict(k="10 / REPRODUCIBILITY",title="Reproduce the registered result",kind="repro",body=[FACTS['repository_url'], f"Formal result: {FACTS['commits']['formal_result']}", f"Implementation freeze: {FACTS['commits']['implementation_freeze']}", f"Maintained release base: {FACTS['commits']['release_base']} · no new formal run here.", f"Pre-window boundary: {FACTS['commits']['pre_window']}"]),
      dict(k="11 / NEXT",title="Improve evidence before expanding claims",kind="next",body=["Reduce fast-layer unsafe events and improve ambiguity handling.", "Test new compositions under a separate preregistered design.", "Add physical-system validation and operator-centered usability work.", "Current claim: a bounded, inspectable recovery research prototype."]),
    ]

def make_video_specs():
    t=TRACE; d=t['decisions']; obs=d[3].get('observation',{}); amb=d[8].get('observation',{})
    return [
      {"name":"Project context","slide_index":0,"voice":"Reality Residual Agent studies bounded recovery decisions for automation teams in a two-dimensional conveyor simulation."},
      {"name":"Recorded trace overview","trace_indices":[0,1,2],"voice":f"Development scenario {t['fault_family']}. This rendered B3 trace {t['episode_id']} is saved evidence, not a live recording: fourteen decisions, seven picked, five escalated; Meterless replay matches the receipt."},
      {"name":"Evidence and common safe action","trace_indices":[0,1,2,3,4],"voice":f"Early checks pause below six samples. At target three, the four displayed checks, first four of fourteen, pass with six. They share a safe action, so COMMON_ACTION records {d[3]['action']} despite multiple passing models."},
      {"name":"Ambiguity and observation","trace_indices":[5,6,7,8],"voice":f"At 3.1 seconds, observed x is {amb['observed_x']:.3f}; nominal x is {amb['nominal_x']:.3f}. The four shown checks, first four of fourteen, pass at thirty-two samples. No common safe action spans the passing hypotheses, so AMBIGUOUS triggers PAUSE_AND_REOBSERVE. The gate need not choose one causal model."},
      {"name":"Hard fault and replay","trace_indices":[9,10,11,12,13],"voice":f"At 3.2 seconds, the recorded policy fault label {d[9]['faults'][0]} latches; policy escalates to a human and remaining targets also escalate. The episode records seven picked, five escalated. Meterless replay matches fourteen decisions and the receipt: persistence evidence, not recovery success."},
      {"name":"Bounded workflow","slide_index":2,"voice":"The fast layer handles residuals; B3 exhaustively searches a bounded DSL. B4 Nemotron proposes or ranks structures. Deterministic fitting, counterfactual checks, the common-action gate, and bounded authorization decide what proceeds. The model never controls actuators."},
      {"name":"Zetaris federation","slide_index":4,"voice":"Zetaris federates evidence: PostgreSQL telemetry plus read-only REST configuration and operations context. The adapter assembles the shared whitelisted C1 packet. Its recorded receipt shows zero T3 and T4 mismatches and identical T8 bytes; this is not a live query."},
      {"name":"Meterless persistence","slide_index":5,"voice":"Meterless uses a pinned local H-MEM reference. The wrapper stores episode runs and manifests, isolates reads, then replays for comparison. Existing checks: four episodes, four replay matches, zero cross-episode reads. These are not eval-v3 scores."},
      {"name":"Formal outcomes","slide_index":6,"voice":"Each condition covers one hundred episodes and six hundred fourteen recoverable challenges. B2: one hundred four recovered, forty-seven unsafe; B3: forty-one, fifteen. STRICT equals B3 by construction. TIEBREAK: thirty-nine recovered, fifty-eight to fifty-nine unsafe, including forty-three to forty-four S2-authorized events. No recovery gain."},
      {"name":"Units and safeguards","slide_index":7,"voice":"Unsafe counts are events. B3: fifteen fast-layer, zero S2-authorized unsafe, one hundred fifty-nine authorizations. The recorded call counter is 1,519 in live fill and scored replay; 1,523 distinct requests in the original live fill; scored replay added zero new live requests and retained 1,523 ledger records; 2,700 proposer events. Replay zero is replay-only, not model non-use."},
      {"name":"Limits and prior work","slide_index":9,"voice":"This is a two-dimensional simulation. Case nine appeared in development; no unseen-composition claim. Safe-stop and fast-layer limits remain. The pre-window includes V0, H1 runtime and evaluation foundations, evidence work, and preregistration. No physical deployment or business outcome is claimed."},
      {"name":"Reproduction and next steps","slide_index":10,"voice":"The registered result and implementation freeze are shown here. The maintained base is not a new run. Next: reduce fast-layer unsafe events, improve ambiguity handling, and preregister new compositions. This remains a bounded research prototype."},
    ]

class Page:
    """Draw one scene in native PowerPoint and matching vector PDF coordinates."""
    def __init__(self, ppt, pdf): self.ppt=ppt; self.pdf=pdf
    def rect(self,x,y,w,h,color,round=False):
        sh=self.ppt.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(*rgb(color)); sh.line.fill.background()
        self.pdf.setFillColor(HexColor('#'+color));
        if round:self.pdf.roundRect(x*72,(H-y-h)*72,w*72,h*72,8,fill=1,stroke=0)
        else:self.pdf.rect(x*72,(H-y-h)*72,w*72,h*72,fill=1,stroke=0)
    def line(self,x1,y1,x2,y2,color=TEAL,width=2):
        self.pdf.setStrokeColor(HexColor('#'+color)); self.pdf.setLineWidth(width); self.pdf.line(x1*72,(H-y1)*72,x2*72,(H-y2)*72)
        self.ppt.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2)).line.color.rgb=RGBColor(*rgb(color))
    def text(self,x,y,w,h,s,size=16,color=WHITE,bold=False,align="left"):
        box=self.ppt.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
        tf.margin_left=Pt(0); tf.margin_right=Pt(0); tf.margin_top=Pt(0); tf.margin_bottom=Pt(0)
        for i,line in enumerate(str(s).split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.name="Aptos"; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor(*rgb(color)); p.space_after=Pt(3)
            p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        font='RRA-DVSans-Bold' if bold else 'RRA-DVSans'; self.pdf.setFillColor(HexColor('#'+color)); self.pdf.setFont(font,size)
        maxw=w*72; lines=[]
        for paragraph in str(s).split('\n'):
            words=paragraph.split(); row=''
            for word in words:
                candidate=(row+' '+word).strip()
                if row and self.pdf.stringWidth(candidate,font,size)>maxw: lines.append(row); row=word
                else: row=candidate
            lines.append(row)
        for i,line in enumerate(lines):
            xx=x*72
            if align=='center': xx+=(maxw-self.pdf.stringWidth(line,font,size))/2
            elif align=='right': xx+=maxw-self.pdf.stringWidth(line,font,size)
            self.pdf.drawString(xx,(H-y)*72-size-i*size*1.28,line)
    def image(self,path,x,y,w,h):
        self.ppt.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
        self.pdf.drawImage(str(path),x*72,(H-y-h)*72,w*72,h*72,mask='auto')

def trace_observation(d):
    return d.get('observation')

def render_trace_ui(outpath, selected=3):
    width,height=1600,900; im=Image.new('RGB',(width,height),(16,27,43)); dr=ImageDraw.Draw(im)
    f=ImageFont.truetype(FONT,18); fb=ImageFont.truetype(FONT_B,21); fs=ImageFont.truetype(FONT,15); fsmall=ImageFont.truetype(FONT,13)
    def box(x,y,w,h,c,r=12): dr.rounded_rectangle((x,y,x+w,y+h),radius=r,fill='#'+c)
    def txt(x,y,s,font=fs,c=WHITE): dr.text((x,y),str(s),font=font,fill='#'+c)
    dr.rectangle((0,0,width,78),fill='#0B1421')
    txt(38,18,"RRA  /  OFFLINE TRACE VIEWER RENDER",fb,TEAL)
    txt(38,49,"Development evidence · not live recording · not formal eval-v3",fsmall,AMBER)
    box(30,100,1540,90,PANEL)
    txt(52,118,f"{TRACE['episode_id']}   ·   DEV SCENARIO: {TRACE['fault_family']}   ·   B3 DEVELOPMENT TRACE",fb)
    t=TRACE['terminal_outcomes']; txt(52,154,f"{len(TRACE['decisions'])} decisions     {t['picked']} picked     {t['escalated']} escalated     Meterless replay: MATCH",f,MUTED)
    box(30,210,930,650,PANEL); box(980,210,590,650,PANEL)
    txt(54,228,"RECORDED DECISION TIMELINE",fb,TEAL); txt(1005,228,"SELECTED RECORD",fb,TEAL)
    headers=[("#",50),("time",105),("target",205),("action",300),("reason",695)]
    for label,x in headers: txt(x,270,label.upper(),fsmall,MUTED)
    rowh=38; first=302
    for i,d in enumerate(TRACE['decisions']):
        y=first+i*rowh
        if i==selected: dr.rounded_rectangle((46,y-3,944,y+31),radius=7,fill='#2D465C',outline='#35D0C3',width=2)
        elif d['action']=='HUMAN_ESCALATION': dr.rounded_rectangle((46,y-3,944,y+31),radius=7,fill='#352C30')
        elif d['action']=='UPDATE_MODEL_AND_REPLAN': dr.rounded_rectangle((46,y-3,944,y+31),radius=7,fill='#22423F')
        txt(54,y,str(i),fs); txt(110,y,fmt(d['time'],1)+'s',fs); txt(215,y,str(d['target_id']),fs)
        action=d['action']; txt(300,y,action,ImageFont.truetype(FONT_B,13),WHITE)
        txt(695,y,d['reason'],fsmall,AMBER if d['reason'] in ('AMBIGUOUS','HARD_FAULT_LATCH') else MUTED)
    d=TRACE['decisions'][selected]
    txt(1005,273,f"Decision {selected} · target {d['target_id']} · t={fmt(d['time'],1)} s",f,WHITE)
    txt(1005,312,d['action'],ImageFont.truetype(FONT_B,16),TEAL)
    txt(1005,345,"Gate / reason: "+d['reason'],fs,AMBER)
    if d['reason']=='COMMON_ACTION':
        txt(1005,369,"Common safe action across passing checks.",fsmall,TEAL)
    elif d['reason']=='AMBIGUOUS':
        txt(1005,369,"No common safe action across passing checks.",fsmall,AMBER)
    yy=403; txt(1005,yy,"OBSERVED FIELDS",fb); yy+=33
    o=trace_observation(d)
    if o:
        for label,key in (("Observed x","observed_x"),("Nominal x","nominal_x"),("Nominal speed","nominal_speed")):
            if key in o: txt(1010,yy,f"{label}: {o[key]:.3f}",fs); yy+=24
    else: txt(1010,yy,"No observation attached to this decision.",fs,MUTED); yy+=25
    yy+=10; txt(1005,yy,f"VISIBLE FIT CHECKS · FIRST {len(d.get('evidence',[]))} OF {d.get('evidence_candidate_total',0)}",fb); yy+=34
    for e in d.get('evidence',[]):
        state='PASS' if e['passed'] else 'not verified'
        txt(1010,yy,f"{e['hypothesis']}: {state} · n={e['n_usable']}",fsmall,TEAL if e['passed'] else MUTED); yy+=24
        if e.get('violations'): txt(1025,yy,"; ".join(e['violations']),ImageFont.truetype(FONT,11),AMBER); yy+=22
    if d.get('faults'): txt(1010,yy+4,"Recorded policy fault label: "+", ".join(d['faults']),fsmall,AMBER)
    box(1004,781,540,54,PANEL2,8); txt(1018,793,"Persisted receipt SHA-256: "+shortsha(TRACE['receipt_sha256']),fsmall,MUTED)
    im.save(outpath)

def add_base(p,sc,index):
    p.rect(0,0,W,H,NAVY); p.rect(0,0,.14,H,TEAL)
    p.text(.62,.32,11,.28,sc['k'],11,TEAL,True)
    p.text(.62,.76,12,.75,sc['title'],30,WHITE,True)
    p.text(.62,7.08,10,.22,"REALITY RESIDUAL AGENT   /   RESEARCH PROTOTYPE",8,MUTED,True)
    p.text(12.14,7.04,.58,.28,f"{index+1:02d}",10,TEAL,True,"right")

def draw_slide(p,sc,index,trace_img):
    add_base(p,sc,index); k=sc['kind']; body=sc['body']
    if k=='cover':
        p.text(.72,1.85,11.8,.72,sc['sub'],20,MUTED)
        p.rect(.72,2.9,11.9,2.4,PANEL,True); p.text(1.05,3.23,10.9,.42,body[0],20,TEAL,True)
        p.text(1.05,3.91,10.8,.43,"Observe → hypothesize → verify → authorize a bounded action",19,WHITE)
        p.text(1.05,4.62,10.8,.35,"Track 03 · Reasoning Architecture · Open Agent Hackathon",14,AMBER,True)
    elif k=='problem':
        for i,(a,b) in enumerate(zip(body[:3],["observation mismatch","timing mismatch","prediction mismatch"])):
            x=.75+i*4.1; p.rect(x,2.35,3.75,1.3,PANEL,True); p.text(x+.18,2.65,3.4,.35,a.upper(),16,AMBER,True,"center"); p.text(x+.18,3.08,3.4,.3,b,13,MUTED,False,"center")
        p.text(.88,4.22,11.5,1.0,body[3],21,WHITE,True)
        p.text(.9,5.3,11.4,.6,"Users: automation and robotics teams evaluating recovery logic before deployment.",16,MUTED)
    elif k=='architecture':
        labels=["OBSERVATION + HISTORY","FAST RESIDUAL","B3 · EXHAUSTIVE DSL","B4 · NEMOTRON PROPOSE / RANK","DETERMINISTIC FIT + COUNTERFACTUAL","COMMON GATE → BOUNDED AUTHORIZATION"]
        for i,l in enumerate(labels):
            x=.7+(i%3)*4.17; y=2.15+(i//3)*1.32; c=TEAL if i in (0,4,5) else PANEL
            p.rect(x,y,3.65,.78,c,True); p.text(x+.15,y+.21,3.35,.35,l,12,NAVY if c==TEAL else WHITE,True,"center")
            if i%3<2:p.text(x+3.69,y+.17,.4,.4,"›",22,AMBER,True,"center")
        p.rect(.72,5.12,11.88,.78,PANEL2,True); p.text(.92,5.31,11.46,.34,"VERIFICATION IS NOT AUTHORIZATION   ·   RECEIPT AFTER AUTHORIZATION   ·   MODEL DOES NOT CONTROL ACTUATORS",12,AMBER,True,"center")
    elif k=='product':
        p.image(trace_img,.7,1.8,7.55,4.75)
        p.rect(8.55,2.0,4.05,3.9,PANEL,True)
        for i,line in enumerate(body): p.text(8.82,2.12+i*.51,3.52,.46,line,12,WHITE,i in (0,1))
    elif k=='zetaris':
        boxes=[(.6,2.5,2.35,1.12,"POSTGRESQL\nTelemetry tables"),(.6,4.15,2.35,1.12,"READ-ONLY REST\nknown_config · ops_context"),(3.55,3.24,2.35,1.2,"ZETARIS\nFederate sources"),(6.55,3.24,2.55,1.2,"SHARED C1 PACKET\nWhitelist + common fields"),(9.75,3.24,2.85,1.2,"POLICY / LOCAL REF\nPacket equivalence")]
        for i,(x,y,w,h,title) in enumerate(boxes): p.rect(x,y,w,h,PANEL if i<2 else TEAL if i==3 else PANEL2,True); p.text(x+.12,y+.2,w-.24,h-.3,title,14,NAVY if i==3 else WHITE,True,"center")
        p.text(3.0,2.83,.45,.45,"→",22,AMBER,True,"center"); p.text(3.0,4.5,.45,.45,"→",22,AMBER,True,"center"); p.text(6.0,3.55,.45,.45,"→",22,AMBER,True,"center"); p.text(9.18,3.55,.45,.45,"→",22,AMBER,True,"center")
        p.rect(.75,5.45,11.85,.7,"382C22",True); p.text(1.0,5.62,11.35,.34,body[-1],14,AMBER,True,"center")
        p.text(.85,6.35,11.6,.3,"Recorded validated receipt; no live SQL request is shown in this demo.",12,MUTED,False,"center")
    elif k=='meterless':
        labels=["RUN RECORD\n+ manifest","PINNED LOCAL\nH-MEM REFERENCE","EPISODE-LOCAL\nREAD / REFUSAL","REPLAY STORED\nMANIFEST","COMPARE DECISIONS\n+ RECEIPT"]
        for i,l in enumerate(labels):
            x=.55+i*2.58; p.rect(x,2.6,2.25,1.25,TEAL if i==4 else PANEL,True); p.text(x+.1,2.88,2.05,.7,l,13,NAVY if i==4 else WHITE,True,"center")
            if i<4:p.text(x+2.27,3.0,.3,.35,"→",18,AMBER,True,"center")
        p.rect(.75,4.55,11.85,.82,PANEL2,True); p.text(.95,4.78,11.45,.4,body[-1],15,TEAL,True,"center")
        p.text(.9,5.7,11.5,.5,"Local reference implementation · not a hosted Meterless cloud service · development checks, not formal scores",13,MUTED,False,"center")
    elif k=='chart':
        b2,b3,b4=source_fields(); vals=[b2['recovered'],b3['recovered'],b4['strict']['recovered_each'],b4['tiebreak']['recovered_each']]
        labels=['B2','B3','B4 STRICT','B4 TIEBREAK']; unsafe=[b2['unsafe_events'],b3['unsafe_events'],b4['strict']['unsafe_events_each'],f"{b4['tiebreak']['unsafe_events_range'][0]}–{b4['tiebreak']['unsafe_events_range'][1]}"]
        p.text(.72,1.8,6.2,.36,f"Recovered challenges / {FACTS['formal']['recoverable_challenges']}",14,MUTED,True)
        for j,(v,u,l) in enumerate(zip(vals,unsafe,labels)):
            y=2.35+j*.76; p.text(.78,y,.98,.32,l,14,WHITE,True); p.rect(1.85,y+.02,3.65*v/110,.28,AMBER if j==3 else TEAL,True); p.text(5.72,y,1.55,.32,f"{v} · unsafe {u}",11,WHITE)
        p.rect(7.55,1.96,5.0,3.6,PANEL,True); p.text(7.86,2.26,4.35,.65,"More unsafe events; no recovery gain",17,AMBER,True)
        p.text(7.86,3.1,4.3,1.9,"100 episodes per condition\nTIEBREAK recovery: 39 each\nUnsafe events: 58–59 across four text conditions\nS2-authorized unsafe events: 43–44\nSTRICT = B3 by construction",13,WHITE)
        p.text(.78,5.68,11.8,.3,"Unsafe values are event counts, not episode rates.",12,MUTED)
    elif k=='interpret':
        a=FACTS['formal']['accounting']; cards=[("RECORDED llm_calls COUNTER",f"{a['live_fill_provider_llm_calls']:,} in live fill and scored replay"),("ORIGINAL LIVE FILL",f"{a['live_fill_distinct_live_requests']:,} distinct requests"),("SCORED REPLAY",f"{a['scored_replay_new_live_requests']} NEW live requests"),("REPLAY CACHE LEDGER",f"{a['scored_replay_cache_ledger_records']:,} records"),("PROPOSER EVENTS",f"{a['total_proposer_events']:,} total · {a['strict_proposer_events']:,} STRICT + {a['tiebreak_proposer_events']:,} TIEBREAK")]
        for i,(title,val) in enumerate(cards):
            if i<3: x=.62+i*4.22; y=2.12; w=3.86
            else: x=.90+(i-3)*6.18; y=3.28; w=5.58
            p.rect(x,y,w,.90,PANEL,True); p.text(x+.18,y+.09,w-.36,.25,title,10,TEAL,True); p.text(x+.18,y+.45,w-.36,.32,val,13,WHITE,True)
        b3=FACTS['formal']['b3']; p.text(.88,4.62,11.55,.57,f"B3 unsafe events: {b3['unsafe_events']} total, all fast-layer; S2-authorized unsafe events: {b3['unsafe_by_source']['s2_authorized']}. B3 authorizations issued: {b3['authorization_count']}.",13,AMBER,True)
        p.text(.88,5.43,11.55,.76,"Live fill = b4_live.json; scored replay = b4.json. Zero NEW requests applies only to replay; model responses were used. Ledger records and proposer events are separate units.",12,MUTED)
    elif k=='sponsors':
        p.rect(.8,2.05,7.2,2.18,PANEL,True); p.text(1.08,2.35,6.65,.4,"NVIDIA NEMOTRON · B4 PROPOSAL / RANKING",15,TEAL,True)
        p.text(1.08,2.95,6.55,.9,"Supplies candidate structure suggestions. Deterministic fit, counterfactual checks, the common-action gate, and explicit bounded authorization decide what proceeds.",14,WHITE)
        p.rect(8.35,2.05,4.15,2.18,PANEL2,True); p.text(8.62,2.35,3.6,.35,"TECHNICAL STACK",14,AMBER,True)
        p.text(8.62,2.9,3.62,1.0,"Python 3.12\nNumPy 2.5.4\nTypeScript H-MEM bridge\nDocker-backed source fixtures",13,WHITE)
        p.text(.9,4.8,11.6,.6,body[-1],15,MUTED,True)
    elif k=='limits':
        for i,t in enumerate(body): p.rect(.82,2.05+i*.87,.2,.2,AMBER,True); p.text(1.18,1.99+i*.87,11.25,.68,t,14,WHITE,i==0)
        p.rect(.82,5.76,11.7,.55,PANEL2,True); p.text(1.0,5.9,11.3,.25,"PRE-EXISTING: HACKATHON_BASELINE.md + pre-window commit; H1 runtime/eval/evidence foundation also predates the window.",11,MUTED,True)
    elif k=='repro':
        for i,t in enumerate(body): p.rect(.82,2.02+i*.78,.18,.18,TEAL,True); p.text(1.16,1.96+i*.78,11.1,.6,t,13,WHITE,i==0)
        p.text(.9,6.2,11.5,.35,"Maintained main is the source base, not a new formal evaluation.",13,AMBER,True)
    elif k=='next':
        for i,t in enumerate(body):
            x=.75+(i%2)*6.1; y=2.05+(i//2)*1.35; p.rect(x,y,5.65,1.0,PANEL,True); p.text(x+.22,y+.23,5.2,.56,t,14,WHITE,i==0)
    return

def build_deck():
    global SLIDES
    SLIDES=make_slide_specs(); OUT.mkdir(parents=True,exist_ok=True)
    prev=OUT/'slide_previews'
    if prev.exists(): shutil.rmtree(prev)
    prev.mkdir()
    trace_img=OUT/'trace_product_view.png'; render_trace_ui(trace_img,selected=3)
    prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
    pdf=canvas.Canvas(str(OUT/'RRA_Submission.pdf'),pagesize=(W*72,H*72))
    for i,sc in enumerate(SLIDES):
        s=prs.slides.add_slide(prs.slide_layouts[6]); p=Page(s,pdf); draw_slide(p,sc,i,trace_img); pdf.showPage()
    prs.save(OUT/'RRA_Submission.pptx'); pdf.save()
    doc=fitz.open(OUT/'RRA_Submission.pdf')
    for i,page in enumerate(doc): page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(prev/f"slide-{i+1:02d}.png")
    make_results_chart()

def make_results_chart():
    b2,b3,b4=source_fields(); vals=[b2['recovered'],b3['recovered'],b4['strict']['recovered_each'],b4['tiebreak']['recovered_each']]
    us=[str(b2['unsafe_events']),str(b3['unsafe_events']),str(b4['strict']['unsafe_events_each']),f"{b4['tiebreak']['unsafe_events_range'][0]}–{b4['tiebreak']['unsafe_events_range'][1]}"]
    labels=['B2','B3','B4 STRICT','B4 TIEBREAK']; fig,ax=plt.subplots(figsize=(11,5.2),facecolor='#'+NAVY); ax.set_facecolor('#'+PANEL)
    bars=ax.bar(range(4),vals,color=['#'+TEAL]*3+['#'+AMBER],width=.58)
    ax.set_ylim(0,125); ax.set_ylabel(f"Recovered challenges / {FACTS['formal']['recoverable_challenges']}",color='#'+WHITE)
    ax.set_xticks(range(4),labels,color='#'+WHITE); ax.tick_params(axis='y',colors='#'+WHITE); ax.spines[:].set_color('#52677c'); ax.grid(axis='y',alpha=.2,color='white')
    for bar,v,u in zip(bars,vals,us): ax.text(bar.get_x()+bar.get_width()/2,v+4,f"{v} recovered\n{u} unsafe events",ha='center',va='bottom',color='#'+WHITE,fontsize=10)
    fig.suptitle("TIEBREAK showed no recovery gain",color='#'+WHITE,fontsize=17,fontweight='bold')
    fig.text(.5,.012,"100 episodes per condition · TIEBREAK 58–59 unsafe / 43–44 S2-authorized across four text conditions · STRICT = B3 by construction",ha='center',color='#'+MUTED,fontsize=8.5)
    fig.tight_layout(rect=[0,.065,1,.92]); fig.savefig(OUT/'Formal_Results.png',dpi=180,facecolor=fig.get_facecolor()); fig.savefig(OUT/'Formal_Results.svg',facecolor=fig.get_facecolor()); plt.close(fig)

def make_viewer():
    payload=json.dumps([TRACE],ensure_ascii=False,separators=(',',':')).replace('</','<\\/')
    template='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>RRA offline development trace</title><style>
body{margin:0;background:#101b2b;color:#f4f7fb;font:15px system-ui,sans-serif}main{max-width:1220px;margin:22px auto;padding:20px}.tag{color:#35d0c3;font-weight:700;letter-spacing:.05em}h1{margin:.3em 0}.warn{padding:12px;background:#382c22;color:#ffb454;border-radius:8px}.panel{padding:16px;background:#17283b;border-radius:12px;margin:14px 0}.toolbar{display:flex;align-items:center;gap:8px;flex-wrap:wrap}button{padding:9px 14px;background:#243349;color:white;border:1px solid #52677c;border-radius:7px}input{flex:1;min-width:180px;accent-color:#35d0c3}.grid{display:grid;grid-template-columns:1.2fr 1fr;gap:16px}.row{display:grid;grid-template-columns:45px 85px 78px 1.4fr 1fr;gap:8px;padding:8px 6px;border-bottom:1px solid #304356;font-size:12px}.row.active{outline:2px solid #35d0c3;background:#2d465c}.head{color:#a9b9c8;font-size:11px}.small{font-size:12px;color:#a9b9c8}.teal{color:#35d0c3}.amber{color:#ffb454}@media(max-width:850px){.grid{grid-template-columns:1fr}.row{grid-template-columns:35px 65px 60px 1fr}}</style><main><div class="tag">OFFLINE SUBMISSION TOOL · DEVELOPMENT TRACE</div><h1>Persisted recovery decisions</h1><p class="warn">Rendered from B3 development evidence. This viewer was produced for this submission; it is not a deployed product interface. This trace is not a registered eval-v3 score.</p><section class="panel"><div class="toolbar"><button id="play">Play</button><button id="step">Step</button><button id="pause">Pause</button><input id="range" type="range" min="0" max="0" value="0"><span id="pos"></span></div><div id="meta" class="small"></div></section><div class="grid"><section class="panel"><div class="tag">ACTUAL DECISION TIMELINE</div><div id="table"></div></section><section class="panel"><div class="tag">SELECTED RECORD</div><div id="detail"></div></section></div><p class="small">All data is embedded. No network calls. Unavailable values are omitted.</p></main><script>const episodes=PAYLOAD;const ep=episodes[0],range=document.querySelector('#range'),table=document.querySelector('#table'),detail=document.querySelector('#detail'),pos=document.querySelector('#pos');let timer=null;const esc=x=>String(x??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));function render(){const i=Number(range.value),d=ep.decisions[i];range.max=ep.decisions.length-1;pos.textContent=(i+1)+' / '+ep.decisions.length;document.querySelector('#meta').textContent=ep.label+' · '+ep.episode_id+' · development scenario '+ep.fault_family+' · source '+ep.source_commit+' · '+ep.terminal_outcomes.picked+' picked / '+ep.terminal_outcomes.escalated+' escalated · replay '+(ep.replay.receipt_match?'MATCH':'MISMATCH');table.innerHTML='<div class="row head"><span>#</span><span>time</span><span>target</span><span>action</span><span>reason</span></div>'+ep.decisions.map((x,n)=>'<div class="row '+(n===i?'active':'')+'"><span>'+n+'</span><span>'+esc(x.time)+' s</span><span>'+esc(x.target_id)+'</span><span>'+esc(x.action)+'</span><span>'+esc(x.reason)+'</span></div>').join('');let o=d.observation;detail.innerHTML='<p class="teal">'+esc(d.action)+'</p><p>t='+esc(d.time)+' s · target '+esc(d.target_id)+' · '+esc(d.reason)+'</p>'+(o?'<p>Observed x '+esc(o.observed_x)+' · nominal x '+esc(o.nominal_x)+' · speed '+esc(o.nominal_speed)+'</p>':'<p class="small">No observation attached to this decision.</p>')+((d.reason==='COMMON_ACTION')?'<p class="small teal">Common safe action across passing checks.</p>':(d.reason==='AMBIGUOUS')?'<p class="small amber">No common safe action across passing checks.</p>':'')+'<h3>Visible fit checks · first '+d.evidence.length+' of '+d.evidence_candidate_total+'</h3>'+d.evidence.map(x=>'<p>'+esc(x.hypothesis)+' · '+(x.passed?'PASS':'not verified')+' · n='+esc(x.n_usable)+(x.violations.length?' · '+esc(x.violations.join('; ')):'')+'</p>').join('')+(d.faults.length?'<p class="amber">Recorded policy fault label: '+esc(d.faults.join(', '))+'</p>':'')+'<p class="small">First '+d.evidence.length+' of '+d.evidence_candidate_total+' candidate checks shown</p>'}document.querySelector('#step').onclick=()=>{range.value=Math.min(Number(range.max),Number(range.value)+1);render()};document.querySelector('#play').onclick=()=>{clearInterval(timer);timer=setInterval(()=>{if(Number(range.value)>=Number(range.max)){clearInterval(timer);return}range.value=Number(range.value)+1;render()},800)};document.querySelector('#pause').onclick=()=>clearInterval(timer);range.oninput=render;render();</script></html>'''.replace('PAYLOAD',payload)
    path=HERE/'demo_viewer.html'; path.write_text(template)
    shutil.copy2(path,OUT/'demo_viewer.html')

def chunks(s,limit=12):
    words=s.split(); return [' '.join(words[i:i+limit]) for i in range(0,len(words),limit)]
def hms(seconds):
    ms=round(seconds*1000); return f"{ms//3600000:02}:{ms//60000%60:02}:{ms//1000%60:02},{ms%1000:03}"

def make_trace_frame(selected,caption):
    path=OUT/'build_only'/'frame_tmp.png'; render_trace_ui(path,selected)
    im=Image.open(path).convert('RGB').resize((1280,720)); return add_caption(im,caption)
def add_caption(im,caption):
    dr=ImageDraw.Draw(im); dr.rectangle((0,630,1280,720),fill=(8,15,24)); font=ImageFont.truetype(FONT,20)
    for j,line in enumerate(textwrap.wrap(caption,78)[:3]): dr.text((42,642+j*25),line,font=font,fill=(245,247,251))
    return im

def make_video():
    global VIDEO_SCENES
    VIDEO_SCENES=make_video_specs(); buildonly=OUT/'build_only'; voice_dir=buildonly/'voice'; aud_dir=buildonly/'audio_segments'
    voice_dir.mkdir(parents=True,exist_ok=True); aud_dir.mkdir(parents=True,exist_ok=True)
    segments=[]; timings=[]; srt=[]; transcript=[]; zh=[]; elapsed=0.0; trace_vis=0.0
    zh_text=[
     "Reality Residual Agent 面向自动化团队，研究二维输送带仿真中的有界恢复决策。",
     "开发场景标签是 SENSOR_STUCK。这是保存的 B3 开发轨迹 ep-014，不是实时录屏：共十四条决策，七个 picked、五个 escalated；Meterless 重放回执匹配。",
     "早期检查因样本少于六条而暂停。target 3 时，显示的四项检查是十四个候选中的前四项，均以六条样本通过。它们支持共同安全动作，因此即使有多个通过模型，COMMON_ACTION 仍记录 UPDATE_MODEL_AND_REPLAN。",
     "三点一秒时，observed x 与 nominal x 如画面所示。十四项中的前四个显示检查均以三十二条样本通过。通过假设之间没有共同安全动作，因此原因是 AMBIGUOUS，动作是 PAUSE_AND_REOBSERVE；门控并不要求唯一因果模型。",
     "三点二秒时，记录中的策略故障标签 BELT_STOP_OR_MOTION_LOSS 锁存，策略向人工升级，其余目标也升级。终态为七个 picked、五个 escalated。Meterless 重放匹配十四条决策和回执；这是持久化一致性证据，不代表恢复成功。",
     "快速层处理即时残差；B3 穷举有限 DSL。B4 Nemotron 提议或排序结构。确定性拟合、反事实检查、共同动作门控与有界授权决定后续步骤。回执记录结果，模型不控制执行器。",
     "Zetaris 联邦证据来源：PostgreSQL 提供白名单观测、出现事件和命令反馈；只读 REST 提供配置与运行上下文。适配器组装共享 C1 packet。归档回执显示 T3、T4 不匹配为零，T8 参考字节相同；没有发起实时查询。",
     "Meterless 使用固定版本的本地 H-MEM reference，而非托管服务。包装层保存 episode 记录和 manifest，隔离 episode 读取，再重放比较。已有开发检查为四个 episode、四次匹配、零跨 episode 读取，不是 eval-v3 分数。",
     "每种正式条件有一百个 episode、六百一十四个可恢复挑战。B2 恢复一百零四个，不安全事件四十七；B3 为四十一个和十五个。STRICT 按构造等于 B3。TIEBREAK 恢复三十九个，不安全事件五十八至五十九个，未测得恢复收益。",
     "不安全计数是事件数。B3 有十五个快速层不安全事件、零个 S2 授权不安全事件，并发出一百五十九次授权。live fill 和 scored replay 文件的 llm_calls 计数字段均为一千五百一十九；这不表示 replay 新发起这些 HTTP 调用。原始 live fill 有一千五百二十三个 distinct live requests；scored replay 新增 live requests 为零，并有一千五百二十三条 cache ledger 记录。Proposer 事件二千七百次。replay 零值不表示没有使用模型。",
     "这是二维仿真。Case 9 曾在开发中出现，不声称未见组合泛化。安全停机与快速层限制仍在。窗口前包含 V0、H1 运行时和评估基础、证据工作及预注册。无实体部署、客户结果或商业回报声明。",
     "正式结果和实现冻结提交号已显示在页面上。维护版不是新运行。后续降低快速层不安全事件、改进含糊处理，并以独立预注册测试新组合。结论仍是有界研究原型。"
    ]
    for i,scene in enumerate(VIDEO_SCENES):
        raw=voice_dir/f"scene-{i+1:02d}.wav"; clip=aud_dir/f"scene-{i+1:02d}.wav"
        txt=scene['voice'].replace("'", " ").replace(":", ".")
        run(['/usr/bin/ffmpeg','-y','-v','error','-f','lavfi','-i',f"flite=text='{txt}':voice=slt",'-ac','2','-ar','44100',str(raw)])
        with wave.open(str(raw),'rb') as wf: speech=wf.getnframes()/wf.getframerate()
        margin=.70; visual=speech+margin
        # Pad only after the complete raw speech. There is deliberately no atrim.
        run(['/usr/bin/ffmpeg','-y','-v','error','-i',str(raw),'-af',f'apad=pad_dur={margin}',str(clip)])
        with wave.open(str(clip),'rb') as wf: clipdur=wf.getnframes()/wf.getframerate()
        if clipdur+0.02 < speech: raise RuntimeError(f"speech truncation in scene {i+1}: {speech} -> {clipdur}")
        pieces=chunks(scene['voice']); weight=[len(x.split()) for x in pieces]; total=sum(weight); cursor=0.0; cue_bounds=[]
        for part,nwords in zip(pieces,weight):
            dur=visual*nwords/total; cue_bounds.append((cursor,cursor+dur,part)); srt.append(f"{len(srt)+1}\n{hms(elapsed+cursor)} --> {hms(elapsed+cursor+dur)}\n{part}\n"); cursor+=dur
        transcript.append(scene['voice']); zh.append(zh_text[i])
        timings.append({'scene':i+1,'name':scene['name'],'speech_seconds':round(speech,6),'clip_seconds':round(clipdur,6),'visual_seconds':round(visual,6),'end_margin_seconds':margin,'truncated_seconds':0,'voice_text':scene['voice'],'raw_voice':f"build_only/voice/{raw.name}"})
        if 'trace_indices' in scene: trace_vis+=visual
        scene['cue_bounds']=cue_bounds; scene['speech_seconds']=speech; scene['visual_seconds']=visual; scene['elapsed']=elapsed
        segments.append(clip); elapsed+=visual
    if not 180<=elapsed<240: raise RuntimeError(f"narrated video duration {elapsed:.2f}s is outside 180–240s; adjust narration content, never cut speech")
    if trace_vis<40: raise RuntimeError(f"trace walkthrough is only {trace_vis:.2f}s, needs at least 40s")
    (OUT/'Transcript_EN.txt').write_text('\n\n'.join(transcript)+'\n')
    (OUT/'Transcript_ZH.txt').write_text('\n\n'.join(zh)+'\n')
    (OUT/'RRA_Demo.srt').write_text('\n'.join(srt))
    (OUT/'Demo_Script.md').write_text('# Narrated demo script\n\n'+'\n\n'.join(f"## {i+1:02d} · {v['name']} ({v['visual_seconds']:.1f}s)\n\n{v['voice_text']}" for i,v in enumerate(timings))+'\n')
    timing={'schema':'rra.video_audio_timing.v2','duration_seconds':round(elapsed,6),'trace_walkthrough_seconds':round(trace_vis,6),'voice':'offline flite slt; no voice cloning','filter':'only end padding; no tempo change or trim','scenes':timings}
    (buildonly/'scene_audio_timing.json').write_text(json.dumps(timing,indent=2,ensure_ascii=False)+'\n')
    # Render one UI frame per actual decision transition; subtitle changes are timed to speech cues.
    frame_files=[]; frame_durs=[]
    for sc in VIDEO_SCENES:
        bounds=set([0.0,sc['visual_seconds']])
        for a,b,_ in sc['cue_bounds']: bounds.add(a); bounds.add(b)
        ids=sc.get('trace_indices',[])
        if ids:
            for j in range(1,len(ids)): bounds.add(sc['visual_seconds']*j/len(ids))
        ts=sorted(bounds)
        for a,b in zip(ts,ts[1:]):
            if b-a<.01: continue
            mid=(a+b)/2; cap=next((c for x,y,c in sc['cue_bounds'] if x<=mid<y),sc['cue_bounds'][-1][2])
            if ids:
                k=min(int(mid/sc['visual_seconds']*len(ids)),len(ids)-1); im=make_trace_frame(ids[k],cap)
            else:
                page=fitz.open(OUT/'RRA_Submission.pdf')[sc['slide_index']]
                pix=page.get_pixmap(matrix=fitz.Matrix(1280/page.rect.width,1280/page.rect.width),alpha=False)
                im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples).resize((1280,720)); im=add_caption(im,cap)
            fn=buildonly/f"vframe-{len(frame_files):04d}.png"; im.save(fn); frame_files.append(fn); frame_durs.append(b-a)
    # Preserve auditable source frames while deleting only intermediate video frames after mux.
    preview_indices=[0,len(frame_files)//3,2*len(frame_files)//3,len(frame_files)-1]
    for i,n in enumerate(preview_indices): shutil.copy2(frame_files[n],OUT/f"video_frame-{i+1:02d}.png")
    make_contact_sheet([OUT/f"video_frame-{i+1:02d}.png" for i in range(4)],OUT/'video_contact_sheet.png')
    audlist=buildonly/'audio_list.txt'; audlist.write_text('\n'.join(f"file '{p}'" for p in segments))
    audio=buildonly/'narration.wav'; run(['/usr/bin/ffmpeg','-y','-v','error','-f','concat','-safe','0','-i',str(audlist),'-c:a','pcm_s16le',str(audio)])
    framelist=buildonly/'frame_list.txt'; rows=[]
    for f,dur in zip(frame_files,frame_durs): rows.extend([f"file '{f}'",f"duration {dur:.6f}"])
    rows.append(f"file '{frame_files[-1]}'"); framelist.write_text('\n'.join(rows))
    run(['/usr/bin/ffmpeg','-y','-v','error','-f','concat','-safe','0','-i',str(framelist),'-i',str(audio),'-t',f'{elapsed:.6f}','-r','30','-c:v','libx264','-pix_fmt','yuv420p','-preset','medium','-crf','22','-c:a','aac','-b:a','160k','-movflags','+faststart',str(OUT/'RRA_Demo.mp4')])
    for f in frame_files:
        try:f.unlink()
        except FileNotFoundError:pass
    (OUT/'Narration_Note.txt').write_text('English narration uses the offline generic flite voice. Raw full-length WAVs and per-scene timings are retained under build_only for review and excluded from the ZIP. Captions are burned in and supplied as SRT. No model API or voice cloning was used.')

def make_contact_sheet(paths,out):
    thumb=Image.new('RGB',(1280,760),(16,27,43)); d=ImageDraw.Draw(thumb); f=ImageFont.truetype(FONT_B,18)
    for i,p in enumerate(paths):
        x=10+(i%2)*638; y=10+(i//2)*372; im=Image.open(p).convert('RGB').resize((620,349)); thumb.paste(im,(x,y)); d.text((x,y+350),f"Video frame {i+1}",font=f,fill='#'+TEAL)
    thumb.save(out)

def make_material_texts():
    # This file is mirrored as a top-level payload file, not buried in source/.
    shutil.copy2(HERE/'START_HERE_ZH.md',OUT/'START_HERE_ZH.md')
    for name in ('Submission_Text_EN.md','Submission_Text_ZH.md','Integration_Explanation.md','Limitations_and_Disclosure.md','Evidence_Source_Index.md','Requirement_Matrix.md'):
        shutil.copy2(HERE/name,OUT/name)
    # Reviewed README/runbooks are copied from the pinned release source_docs bundle.
    shutil.copy2(HERE/'source_docs/README.md',OUT/'README.md')
    shutil.copy2(HERE/'source_docs/README.en.md',OUT/'README.en.md')
    (OUT/'docs').mkdir(exist_ok=True)
    shutil.copy2(HERE/'source_docs/docs/DEMO_RUNBOOK.md',OUT/'docs/DEMO_RUNBOOK.md')
    shutil.copy2(HERE/'source_docs/docs/REPRODUCING_EVAL_V3.md',OUT/'docs/REPRODUCING_EVAL_V3.md')

def package():
    make_material_texts()
    archive=OUT/'RRA_Submission_Package.zip'; entries={}
    def put(name,path): entries[name]=Path(path).read_bytes()
    for name in ('RRA_Submission.pptx','RRA_Submission.pdf','RRA_Demo.mp4','RRA_Demo.srt','Transcript_EN.txt','Transcript_ZH.txt','Demo_Script.md','Narration_Note.txt','demo_viewer.html','Formal_Results.png','Formal_Results.svg','trace_product_view.png','video_contact_sheet.png'):
        put(name,OUT/name)
    for p in sorted((OUT/'slide_previews').glob('*.png')): put('slide_previews/'+p.name,p)
    for p in sorted(OUT.glob('video_frame-*.png')): put(p.name,p)
    for name in ('START_HERE_ZH.md','README.md','README.en.md','Submission_Text_EN.md','Submission_Text_ZH.md','Integration_Explanation.md','Limitations_and_Disclosure.md','Evidence_Source_Index.md','Requirement_Matrix.md'): put(name,OUT/name)
    for name in ('DEMO_RUNBOOK.md','REPRODUCING_EVAL_V3.md'): put('docs/'+name,OUT/'docs'/name)
    support=['build_submission.py','validate_viewer.mjs','requirements-material.txt','demo_viewer.html','Submission_Text_EN.md','Submission_Text_ZH.md','Integration_Explanation.md','Limitations_and_Disclosure.md','Evidence_Source_Index.md','Requirement_Matrix.md','START_HERE_ZH.md']
    for name in support: put('source/'+name,HERE/name)
    for p in sorted((HERE/'data').glob('*')):
        if p.is_file(): put('source/data/'+p.name,p)
    for p in sorted((HERE/'source_docs').rglob('*')):
        if p.is_file(): put('source/source_docs/'+str(p.relative_to(HERE/'source_docs')),p)
    payload_manifest=''.join(f"{hashlib.sha256(data).hexdigest()}  {name}\n" for name,data in sorted(entries.items()))
    entries['PAYLOAD_SHA256SUMS.txt']=payload_manifest.encode()
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=8) as z:
        for name,data in entries.items(): z.writestr(name,data)
    # External checksums are separate. Build-only voice WAVs are evidence, not exports.
    files=[p for p in OUT.rglob('*') if p.is_file() and p.name!='SHA256SUMS.txt' and 'build_only' not in p.parts]
    (OUT/'SHA256SUMS.txt').write_text(''.join(f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.relative_to(OUT)}\n" for p in sorted(files)))

def main():
    global OUT
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--output-dir',type=Path,default=HERE/'rebuild-output',help='relative or absolute export directory; default is adjacent to this source')
    args=ap.parse_args(); OUT=args.output_dir.expanduser().resolve(); OUT.mkdir(parents=True,exist_ok=True)
    if not Path('/usr/bin/ffmpeg').exists() or not Path('/usr/bin/ffprobe').exists(): raise SystemExit('ffmpeg and ffprobe are required')
    if not Path(FONT).exists() or not Path(FONT_B).exists(): raise SystemExit('DejaVu Sans fonts are required; see START_HERE_ZH.md')
    make_viewer(); build_deck(); make_video(); package()
    print(f"Built {OUT}")

if __name__=='__main__': main()
